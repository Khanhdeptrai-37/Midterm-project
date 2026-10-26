"""
EE136IU Capstone Mini-Project - Topic 1: Comprehensive Model Benchmark
Automated Presentation Generator:
1. Generates 'EE136IU_Topic1_Classical_ML_Presentation.pptx' (Dedicated Classical ML deck from ML.ipynb)
2. Generates 'EE136IU_Topic1_Complete_Benchmark_Presentation.pptx' (Consolidated 20-slide Master deck)
Adheres strictly to the course Clean White Academic design system (#FFFFFF canvas, #F8FAFC cards, #0284C7 accents, #0F172A text).
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Design System Colors
BG_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BG      = RGBColor(0xF8, 0xFA, 0xFC)
CARD_BORDER  = RGBColor(0xE2, 0xE8, 0xF0)
CARD_BG_ALT  = RGBColor(0xFA, 0xFA, 0xF9)
TEXT_MAIN    = RGBColor(0x0F, 0x17, 0x2A)
TEXT_MUTED   = RGBColor(0x64, 0x74, 0x8B)
TEXT_DARK    = RGBColor(0x1E, 0x29, 0x3B)
TEXT_SLATE   = RGBColor(0x33, 0x41, 0x55)
TECH_BLUE    = RGBColor(0x02, 0x84, 0xC7)
ACCENT_BLUE  = RGBColor(0x25, 0x63, 0xEB)
TIER1_GREEN  = RGBColor(0x10, 0xB9, 0x81)
TIER2_BLUE   = RGBColor(0x3B, 0x82, 0xF6)
TIER3_PINK   = RGBColor(0xEC, 0x48, 0x99)
AMBER_WARN   = RGBColor(0xD9, 0x77, 0x06)
PURPLE_STAT  = RGBColor(0x8B, 0x5C, 0xF6)

FONT_TITLE = 'Manrope'
FONT_BODY  = 'Inter'
FONT_CODE  = 'Consolas'

def set_shape_flat(shape, fill_color=CARD_BG, line_color=CARD_BORDER, line_width=1.5):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = line_color
    shape.line.width = Pt(line_width)

def add_header(slide, tracker_text, title_text):
    # Tracker Pill / Breadcrumb
    tx_trk = slide.shapes.add_textbox(Inches(0.80), Inches(0.32), Inches(11.70), Inches(0.30))
    tf_trk = tx_trk.text_frame
    tf_trk.word_wrap = True
    p_trk = tf_trk.paragraphs[0]
    p_trk.text = tracker_text
    p_trk.font.name = FONT_TITLE
    p_trk.font.size = Pt(11)
    p_trk.font.bold = True
    p_trk.font.color.rgb = TECH_BLUE
    
    # Title
    tx_ttl = slide.shapes.add_textbox(Inches(0.80), Inches(0.60), Inches(11.70), Inches(0.55))
    tf_ttl = tx_ttl.text_frame
    tf_ttl.word_wrap = True
    p_ttl = tf_ttl.paragraphs[0]
    p_ttl.text = title_text
    p_ttl.font.name = FONT_TITLE
    p_ttl.font.size = Pt(21)
    p_ttl.font.bold = True
    p_ttl.font.color.rgb = TEXT_MAIN

def add_footer(slide, current_slide, total_slides, course_sub="Topic 1: Comprehensive Model Benchmark"):
    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), Inches(6.85), Inches(11.73), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.color.rgb = CARD_BORDER
    line.line.width = Pt(0.5)
    
    # Footer Left
    tx_fl = slide.shapes.add_textbox(Inches(0.80), Inches(6.95), Inches(9.00), Inches(0.30))
    tf_fl = tx_fl.text_frame
    p_fl = tf_fl.paragraphs[0]
    p_fl.text = f"EE136: Machine Learning for Wireless Communications | {course_sub}"
    p_fl.font.name = FONT_BODY
    p_fl.font.size = Pt(9.5)
    p_fl.font.color.rgb = TEXT_MUTED

    # Footer Right
    tx_fr = slide.shapes.add_textbox(Inches(10.50), Inches(6.95), Inches(2.03), Inches(0.30))
    tf_fr = tx_fr.text_frame
    p_fr = tf_fr.paragraphs[0]
    p_fr.alignment = PP_ALIGN.RIGHT
    p_fr.text = f"Slide {current_slide} of {total_slides}"
    p_fr.font.name = FONT_TITLE
    p_fr.font.size = Pt(9.5)
    p_fr.font.bold = True
    p_fr.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, fill_color=CARD_BG, border_color=CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_shape_flat(card, fill_color, border_color)
    return card

def add_bullet_list(text_frame, items, font_size=12, text_color=TEXT_MAIN, bold_prefix_len=0):
    for idx, item in enumerate(items):
        if idx == 0 and len(text_frame.paragraphs) > 0 and text_frame.paragraphs[0].text == '':
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.space_after = Pt(4)
        p.font.name = FONT_BODY
        p.font.size = Pt(font_size)
        
        if isinstance(item, tuple):
            prefix, body = item
            r1 = p.add_run()
            r1.text = prefix + " "
            r1.font.bold = True
            r1.font.color.rgb = TECH_BLUE
            r2 = p.add_run()
            r2.text = body
            r2.font.bold = False
            r2.font.color.rgb = text_color
        else:
            p.text = item
            p.font.color.rgb = text_color

print("Presentation generator module loaded successfully.")
